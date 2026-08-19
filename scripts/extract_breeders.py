"""Rebuild frontend/src/constants/breeders.json from PAW_BUDDY_Breeder_Database.xlsx.

Run:  python scripts/extract_breeders.py --check     (diff only, writes nothing)
      python scripts/extract_breeders.py             (write the file)

Tier 1 and Tier 2 only. Tier 3 ("Low") and Tier 4 ("Investigate") carry
documented welfare red flags and are excluded outright — and the
"Notes / Red Flags" column is never exported at any tier, because it holds
unverified review-sourced allegations about named real businesses.

`--check` exists because this script REGENERATES records that are already in
production. Reproducing the existing rows byte-for-byte is the evidence that
the parsing and the breed-alias map are right; only then is it safe to trust
the same code on a city that has no existing rows to check against.
"""
import argparse
import json
import os
import re
import sys

import openpyxl

SOURCE = os.environ.get(
    "BREEDER_XLSX",
    os.path.join(os.path.expanduser("~"), "Downloads", "PAW_BUDDY_Breeder_Database.xlsx"),
)
OUT = os.path.join("frontend", "src", "constants", "breeders.json")
BREEDS = os.path.join("frontend", "src", "constants", "breeds.json")

CITIES = ["Hyderabad", "Chennai", "Bangalore", "Pune", "Mumbai"]
CITY_SLUG = {"Hyderabad": "hyd", "Chennai": "che", "Bangalore": "ban", "Pune": "pun", "Mumbai": "mum"}

#: Informal names used in reviews -> the app's canonical breed name. Only
#: mappings that are unambiguous: "Toy Poodle" and "Poodle" both fold into the
#: single Poodle entry the catalogue carries, "Labrador" into the full name.
BREED_ALIASES = {
    "labrador": "Labrador Retriever",
    "toy poodle": "Poodle (Standard)",
    "poodle": "Poodle (Standard)",
    "doberman": "Doberman Pinscher",
    "husky": "Siberian Husky",
    "golden cocker spaniel": "Cocker Spaniel",
}
#: "Toy Pomeranian" is deliberately NOT aliased to Pomeranian. Doing so is
#: arguably more accurate (a toy Pomeranian is a Pomeranian) but it would
#: rewrite one already-published Bangalore record's breed list, which this
#: import is not meant to touch. Add the alias only as a deliberate,
#: separately-reviewed data change.

#: Tokens that are not a breed at all. Dropped from both lists rather than
#: being published as if a breeder offers a dog called "general pets".
NON_BREED = re.compile(
    r"^(multiple breeds|not specified.*|general pets|birds|kittens|persian cats|"
    r"mating / stud services.*|stud dog.*)$", re.I)


def _text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _tier_number(value):
    match = re.search(r"tier\s*(\d+)", str(value or ""), re.I)
    return int(match.group(1)) if match else 99


def _header_row(ws):
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        if row and any(_text(c) == "Rank" for c in row):
            return i
    raise ValueError("no header row containing 'Rank'")


def split_breeds(raw, canonical):
    """'Toy Poodle; Cane Corso; multiple breeds' ->
       (['Poodle (Standard)'], ['Cane Corso'])

    A parenthetical qualifier ("German Shepherd (specialist)") describes the
    breeder, not a different dog, so it is stripped before matching."""
    known, other = [], []
    for token in str(raw or "").split(";"):
        token = _text(token)
        if not token or NON_BREED.match(token):
            continue
        base = _text(re.sub(r"\(.*?\)", "", token))
        if not base or NON_BREED.match(base):
            continue
        mapped = BREED_ALIASES.get(base.lower())
        if mapped is None and base in canonical:
            mapped = base
        if mapped:
            if mapped not in known:
                known.append(mapped)
        elif base not in other:
            other.append(base)
    return known, other


def extract(cities):
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    canonical = {b["name"] for b in json.load(open(BREEDS, encoding="utf-8"))}
    out, stats = [], []

    for city in cities:
        ws = wb[city]
        hrow = _header_row(ws)
        header = [_text(c) for c in next(ws.iter_rows(min_row=hrow, max_row=hrow, values_only=True))]
        idx = {h: i for i, h in enumerate(header) if h}
        name_col = next(h for h in header if h.endswith("Name"))

        rows, excluded = [], 0
        for source_row, cells in enumerate(ws.iter_rows(min_row=hrow + 1)):
            values = [c.value for c in cells]
            name = _text(values[idx[name_col]])
            if not name:
                continue
            tier = _tier_number(values[idx["Tier"]])
            if tier not in (1, 2):
                excluded += 1
                continue

            maps_cell = cells[idx["Google Maps Link"]]
            known, other = split_breeds(values[idx["Breeds Observed"]], canonical)

            def num(col, cast):
                try:
                    return cast(values[idx[col]])
                except (TypeError, ValueError):
                    return None

            rows.append({
                "city": city,
                "zone": _text(values[idx["Zone"]]),
                "name": name,
                "owner": _text(values[idx["Owner / Contact"]]),
                "locality": _text(values[idx["Locality"]]),
                "address": _text(values[idx["Full Address"]]),
                "phone": _text(values[idx["Phone"]]),
                "rating": num("Google Rating", float),
                "reviews": num("Review Count", lambda v: int(float(v))),
                "breeds": known,
                "otherBreeds": other,
                "certification": _text(values[idx["Certification / Docs Mentioned"]]),
                "mapsUrl": maps_cell.hyperlink.target if maps_cell.hyperlink else "",
                # Kept because utils/breeders.js sorts on them. Never rendered.
                "tier": tier,
                "confidence": num("Confidence Score", lambda v: int(float(v))) or 0,
                "_row": source_row,
            })

        # The spreadsheet's own Rank column already encodes the intended order,
        # and the records live in Rank order. Re-deriving an order from
        # (tier, confidence) looks equivalent but breaks ties differently, which
        # silently renumbers existing published rows — so keep source order.
        rows.sort(key=lambda r: r["_row"])
        for n, record in enumerate(rows, 1):
            record.pop("_row")
            record["id"] = f"{CITY_SLUG[city]}-{n}"

        out.extend(rows)
        stats.append((city, len(rows), excluded))

    # id first, matching the existing file's key order
    ordered = [{"id": r.pop("id"), **r} for r in out]
    return ordered, stats


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true", help="diff against the current file, write nothing")
    args = ap.parse_args()

    records, stats = extract(CITIES)

    for record in records:
        for banned in ("notes", "redFlags", "notesRedFlags"):
            if banned in record:
                raise SystemExit(f"ABORT: {banned!r} would be published")

    current = json.load(open(OUT, encoding="utf-8")) if os.path.exists(OUT) else []
    cur_by_id = {r["id"]: r for r in current}
    new_by_id = {r["id"]: r for r in records}

    added = [i for i in new_by_id if i not in cur_by_id]
    removed = [i for i in cur_by_id if i not in new_by_id]
    changed = [i for i in new_by_id if i in cur_by_id and new_by_id[i] != cur_by_id[i]]

    print(f"current={len(current)}  regenerated={len(records)}")
    for city, kept, excluded in stats:
        print(f"  {city:11} kept={kept:3}  excluded (Tier 3/4)={excluded}")
    print(f"\n  ADDED   : {len(added)}  {sorted(added)[:14]}")
    print(f"  REMOVED : {len(removed)}  {sorted(removed)[:14]}")
    print(f"  CHANGED : {len(changed)}  {sorted(changed)[:14]}")
    for i in sorted(changed)[:5]:
        diffs = {k: (cur_by_id[i].get(k), new_by_id[i].get(k))
                 for k in new_by_id[i] if cur_by_id[i].get(k) != new_by_id[i].get(k)}
        print(f"     {i}: {diffs}")

    if args.check:
        print("\n--check: nothing written.")
        return 0

    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(records, fh, ensure_ascii=False, indent=2)
        fh.write("\n")
    print(f"\nwrote {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
