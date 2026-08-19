"""Build frontend/src/constants/groomers.json from PAW_BUDDY_Groomer_Database.xlsx.

Run:  python scripts/extract_groomers.py

Mirrors extract_vets.py's two rules, mechanically rather than by remembering:

1. ONLY user-facing columns are written. Tier, Confidence Score, the component
   scores, "Documentation Mentioned", and above all "Notes" are read (tier and
   confidence decide the ORDER) and then dropped — they never reach the JSON,
   so no later UI work can accidentally render them.

2. Only Tier 1 and Tier 2 are exported. Tier 3 is internal-only; the
   spreadsheet remains its home.

The Google Maps column displays the text "Open in Maps" and carries the real
URL as a cell HYPERLINK. Reading .value gives you the label, not the link.

City sheets are auto-discovered (anything that isn't a known summary sheet),
not hardcoded — the workbook currently has only Hyderabad, and new city tabs
should extract with no code change, same as they will for vets once added.
"""
import json
import os
import re
import sys

import openpyxl

SOURCE = os.environ.get(
    "GROOMER_XLSX",
    os.path.join(os.path.expanduser("~"), "Downloads", "PAW_BUDDY_Groomer_Database.xlsx"),
)
OUT = os.path.join("frontend", "src", "constants", "groomers.json")

#: Non-city sheets to skip when discovering which tabs hold data.
NON_CITY_SHEETS = {"Zone Summary", "Scoring Method", "Summary", "Groomer Summary"}

#: Exported verbatim. Anything not listed here is internal by construction.
FIELDS = {
    "name": "Business Name",
    "groomers": "Groomer(s)",
    "locality": "Locality",
    "address": "Full Address",
    "phone": "Phone",
    "rating": "Google Rating",
    "reviews": "Review Count",
    "services": "Services Observed",
    "specialFeatures": "Special Features",
}
#: Read for ordering only, then discarded — never written to the JSON.
SORT_ONLY = {"tier": "Tier", "confidence": "Confidence Score"}


def _text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _tier_number(value):
    """'Tier 2 - Verify Before Contact' -> 2. Unrecognised text sorts last and
    is excluded from export, so a reworded tier can never silently publish."""
    match = re.search(r"tier\s*(\d+)", str(value or ""), re.I)
    return int(match.group(1)) if match else 99


def _header_row(ws):
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), 1):
        if row and any(_text(c) == "Rank" for c in row):
            return i
    raise ValueError("no header row containing 'Rank'")


def _city_slug(city, taken):
    base = re.sub(r"[^a-z]", "", city.lower())[:3] or "cty"
    slug = base
    n = 1
    while slug in taken:
        n += 1
        slug = f"{base}{n}"
    taken.add(slug)
    return slug


def extract():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    cities = [s for s in wb.sheetnames if s not in NON_CITY_SHEETS]
    out, stats, slugs = [], [], set()

    for city in cities:
        ws = wb[city]
        hrow = _header_row(ws)
        header = [_text(c) for c in next(ws.iter_rows(min_row=hrow, max_row=hrow, values_only=True))]
        idx = {h: i for i, h in enumerate(header) if h}

        missing = [c for c in list(FIELDS.values()) + list(SORT_ONLY.values()) if c not in idx]
        if missing:
            raise ValueError(f"{city}: missing column(s) {missing}")

        slug = _city_slug(city, slugs)
        rows, excluded = [], 0
        for cells in ws.iter_rows(min_row=hrow + 1):
            values = [c.value for c in cells]
            name = _text(values[idx["Business Name"]])
            if not name:
                continue

            tier = _tier_number(values[idx["Tier"]])
            if tier not in (1, 2):
                excluded += 1
                continue

            maps_cell = cells[idx["Google Maps Link"]]
            maps_url = maps_cell.hyperlink.target if maps_cell.hyperlink else ""

            record = {key: _text(values[idx[col]]) for key, col in FIELDS.items()}
            record["city"] = city
            record["mapsUrl"] = maps_url or ""

            try:
                record["rating"] = float(record["rating"]) if record["rating"] else None
            except ValueError:
                record["rating"] = None
            try:
                record["reviews"] = int(float(record["reviews"])) if record["reviews"] else None
            except ValueError:
                record["reviews"] = None

            confidence = values[idx["Confidence Score"]]
            try:
                confidence = float(confidence)
            except (TypeError, ValueError):
                confidence = 0.0
            # _sort_tier / _sort_conf are stripped below — exist only to decide
            # order, must not survive into the exported record.
            record["_sort_tier"] = tier
            record["_sort_conf"] = confidence
            rows.append(record)

        rows.sort(key=lambda r: (r["_sort_tier"], -r["_sort_conf"], r["name"].lower()))
        for n, record in enumerate(rows, 1):
            record.pop("_sort_tier")
            record.pop("_sort_conf")
            record["id"] = f"{slug}-groomer-{n}"

        out.extend(rows)
        stats.append((city, len(rows), excluded))

    return out, stats


def main():
    records, stats = extract()

    allowed = set(FIELDS) | {"city", "mapsUrl", "id"}
    for record in records:
        leaked = set(record) - allowed
        if leaked:
            raise SystemExit(f"ABORT: internal field(s) would be published: {sorted(leaked)}")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as fh:
        json.dump(records, fh, ensure_ascii=False, indent=2)
        fh.write("\n")

    print(f"wrote {OUT}")
    for city, kept, excluded in stats:
        print(f"  {city:11} exported={kept:3}  excluded (Tier 3)={excluded}")
    print(f"  {'TOTAL':11} exported={len(records)}")
    print(f"  fields published: {sorted(allowed)}")


if __name__ == "__main__":
    sys.exit(main())
